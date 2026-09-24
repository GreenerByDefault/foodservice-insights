from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Literal, cast

import chardet
import openpyxl
import pandas as pd

from gbd_foodservice_insights_lab.extraction.llm import experimental_csv_extraction_to_markdown

__all__ = [
    "detect_excel_header_row",
    "detect_file_encoding",
    "detect_single_multi_sheet_excel",
    "experimental_excel_extraction_code",
    "extract_dates_from_sheet_filenames",
    "extract_excel_sheet_info",
    "extract_month_year_from_filename",
    "filter_hidden_and_temp_files",
    "get_excel_headers",
    "get_filtered_data_files",
    "read_csv_with_auto_detection",
    "read_in_all_data_files",
]


def filter_hidden_and_temp_files(
    files: list[Path],
    file_type: Literal["csv", "excel", "pdf"] = "csv",
) -> list[Path]:
    """Filter out hidden and temporary files from a list of paths.

    Hidden files start with ``.`` and temporary files start with ``~``. Excel
    temp files additionally start with ``~$``.

    Args:
        files: Candidate file paths to filter.
        file_type: File family the rules apply to (``"csv"``, ``"excel"``, or ``"pdf"``).

    Returns:
        Paths that are not hidden or temporary.
    """
    if file_type == "excel":
        return [f for f in files if not f.name.startswith((".", "~$", "~"))]
    return [f for f in files if not f.name.startswith((".", "~"))]


def get_filtered_data_files(
    path: str | Path,
    file_types: list[Literal["csv", "excel", "pdf"]] | None = None,
    recursive: bool = False,
) -> dict[str, list[Path]]:
    """List data files in a directory, grouped by file family.

    Args:
        path: Directory to search.
        file_types: File families to include; defaults to all three.
        recursive: Whether to search subdirectories.

    Returns:
        Mapping with keys ``"csv"``, ``"excel"``, and ``"pdf"`` and lists of
        matching paths as values.
    """
    base = Path(path)
    globber = base.rglob if recursive else base.glob

    if file_types is None:
        file_types = ["csv", "excel", "pdf"]

    result = {"csv": [], "excel": [], "pdf": []}

    if "csv" in file_types:
        result["csv"] = filter_hidden_and_temp_files(list(globber("*.csv")), "csv")

    if "excel" in file_types:
        excel_raw = list(globber("*.xls")) + list(globber("*.xlsx"))
        result["excel"] = filter_hidden_and_temp_files(excel_raw, "excel")

    if "pdf" in file_types:
        result["pdf"] = filter_hidden_and_temp_files(list(globber("*.pdf")), "pdf")

    return result


def detect_single_multi_sheet_excel(
    path: str | Path,
    min_sheets: int = 2,
) -> tuple[bool, Path | None]:
    """Detect if a directory contains exactly one Excel file with multiple sheets.

    Args:
        path: Directory to inspect.
        min_sheets: Minimum sheet count to qualify as multi-sheet.

    Returns:
        Tuple of (matched flag, the Excel file path if matched else ``None``).
    """
    files = get_filtered_data_files(path, file_types=["excel"])
    excel_files = files["excel"]

    if len(excel_files) != 1:
        return False, None

    excel_file = excel_files[0]

    try:
        wb = openpyxl.load_workbook(excel_file, read_only=True)
        sheet_count = len(wb.sheetnames)
        wb.close()
        if sheet_count >= min_sheets:
            return True, excel_file
        return False, None
    except Exception:
        return False, None


def detect_file_encoding(filepath: str | Path) -> str:
    """Detect a file's text encoding with the chardet library.

    Args:
        filepath: Path of the file to sniff.

    Returns:
        Detected encoding string, falling back to ``"utf-8"`` on error.
    """
    try:
        with open(filepath, "rb") as f:
            raw_data = f.read(100000)
            result = chardet.detect(raw_data)
            return result["encoding"] or "utf-8"
    except Exception:
        return "utf-8"


def get_excel_headers(worksheet: Any, header_row: int = 1) -> list[Any]:
    """Extract header values from an Excel worksheet.

    Args:
        worksheet: An openpyxl worksheet instance.
        header_row: 1-indexed row number to read headers from.

    Returns:
        List of cell values from the header row.
    """
    return [c.value for c in worksheet[header_row]]


def _format_preview_cell(value: Any) -> str:
    """Format a worksheet cell value safely for notebook-style preview output."""
    if value is None:
        return "<blank>"
    text = str(value).strip()
    return text if text else "<blank>"


def _looks_numeric(value: Any) -> bool:
    """Return True when a value is numeric-like and therefore less likely to be a header label."""
    if value is None:
        return False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True
    text = str(value).strip()
    if not text:
        return False
    try:
        float(text.replace(",", ""))
        return True
    except ValueError:
        return False


def detect_excel_header_row(
    worksheet: Any,
    scan_rows: int = 20,
) -> int:
    """Guess the most likely header row in the first part of an Excel worksheet.

    Args:
        worksheet: An openpyxl worksheet instance.
        scan_rows: Number of leading rows to score.

    Returns:
        1-indexed row number that scored highest as a likely header.
    """
    best_row = 1
    best_score = float("-inf")

    worksheet_max_row = getattr(worksheet, "max_row", None)
    max_rows_to_scan = scan_rows if worksheet_max_row is None else min(scan_rows, worksheet_max_row)
    for row_idx in range(1, max_rows_to_scan + 1):
        values = [
            cell
            for cell in next(
                worksheet.iter_rows(
                    min_row=row_idx,
                    max_row=row_idx,
                    values_only=True,
                )
            )
        ]
        nonempty = [value for value in values if value is not None and str(value).strip()]
        if not nonempty:
            continue

        text_like_count = sum(not _looks_numeric(value) for value in nonempty)
        numeric_like_count = len(nonempty) - text_like_count
        unique_text_count = len({str(value).strip().lower() for value in nonempty})
        metadata_penalty = 0
        if len(nonempty) <= 2 and any(":" in str(value) for value in nonempty):
            metadata_penalty += 3
        if len(nonempty) == 1:
            metadata_penalty += 2

        score = (
            len(nonempty) * 3
            + text_like_count * 2
            + unique_text_count
            - numeric_like_count * 2
            - metadata_penalty
        )

        if score > best_score:
            best_score = score
            best_row = row_idx

    return best_row


def extract_excel_sheet_info(filepath: str | Path) -> tuple[dict[str, list], int]:
    """Extract column names and total row counts from all sheets in an Excel file.

    Args:
        filepath: Excel file to inspect.

    Returns:
        Tuple of (mapping of sheet name to header columns, total data-row count
        across all sheets).
    """
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        sheet_columns = {}
        total_rows = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_row = detect_excel_header_row(ws)
            cols = get_excel_headers(ws, header_row=header_row)
            sheet_columns[sheet_name] = cols
            total_rows += max(ws.max_row - header_row, 0)

        wb.close()
        return sheet_columns, total_rows
    except Exception as e:
        print(f"Error reading {Path(filepath).name}: {e}")
        return {}, 0


def read_csv_with_auto_detection(
    filepath: str | Path,
    nrows: int | None = None,
) -> tuple[pd.DataFrame, str, str]:
    """Read a CSV file with automatic encoding and delimiter detection.

    Args:
        filepath: Path of the CSV file to read.
        nrows: Optional row cap, useful for previews.

    Returns:
        Tuple of (loaded DataFrame, detected encoding, detected delimiter).
    """
    encoding = detect_file_encoding(filepath)

    delimiter = ","
    try:
        with open(filepath, encoding=encoding) as f_sniff:
            sample = f_sniff.read(5000)
            dialect = csv.Sniffer().sniff(sample)
            delimiter = dialect.delimiter
    except Exception:
        delimiter = ","

    df = pd.read_csv(filepath, encoding=encoding, delimiter=delimiter, nrows=nrows)
    return df, encoding, delimiter


def extract_month_year_from_filename(
    filename: str | Path,
    date_format: str | None = None,
) -> pd.Timestamp | None:
    """Extract month and year from a filename using common date patterns.

    Args:
        filename: Filename or path to inspect.
        date_format: Optional explicit ``strptime`` format to try first.

    Returns:
        Parsed ``pd.Timestamp`` or ``None`` when no pattern matched.
    """
    filename_str = str(Path(filename).stem)

    if date_format:
        try:
            date_obj = datetime.strptime(filename_str, date_format)
            return cast(pd.Timestamp, pd.Timestamp(date_obj))
        except ValueError:
            pass

    patterns = [
        (r"(\d{4})[-_.](\d{1,2})", "%Y-%m"),
        (r"(\d{1,2})[-_.](\d{4})", "%m-%Y"),
        (
            r"(January|February|March|April|May|June|July|August|September|October|November"
            r"|December)\.?[-_ ]+(\d{4})",
            "%B %Y",
        ),
        (r"(Sept)\.?[-_ ]+(\d{4})", "%b %Y", lambda x: x.replace("Sept", "Sep")),
        (r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\.?[-_ ]+(\d{4})", "%b %Y"),
        (
            r"(January|February|March|April|May|June|July|August|September|October|November"
            r"|December)(\d{4})",
            "%B %Y",
            lambda x: re.sub(r"([a-zA-Z])(\d)", r"\1 \2", x),
        ),
        (
            r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)(\d{4})",
            "%b %Y",
            lambda x: re.sub(r"([a-zA-Z])(\d)", r"\1 \2", x),
        ),
        (r"(\d{4})(\d{2})", "%Y%m"),
    ]

    for pattern_info in patterns:
        if len(pattern_info) == 3:
            pattern, fmt, normalizer = pattern_info
        else:
            pattern, fmt = pattern_info
            normalizer = None

        match = re.search(pattern, filename_str, re.IGNORECASE)
        if match:
            try:
                date_str = match.group(0)
                if normalizer:
                    date_str = normalizer(date_str)
                date_str = re.sub(r"\.", "", date_str)
                date_str = re.sub(r"[-_]+", " ", date_str)
                date_str = " ".join(date_str.split())
                date_obj = datetime.strptime(date_str, fmt)
                return cast(pd.Timestamp, pd.Timestamp(date_obj))
            except ValueError, AttributeError:
                continue

    return None


def extract_dates_from_sheet_filenames(
    df: pd.DataFrame,
    filename_column: str = "sheet_filename",
    default_year: int | None = None,
    date_format: str | None = None,
    show_warnings: bool = True,
) -> pd.DataFrame:
    """Extract dates from sheet or file names and add a ``date`` column.

    Args:
        df: DataFrame containing the filename column.
        filename_column: Name of the column holding filenames.
        default_year: Year to use when filenames contain only a month.
        date_format: Optional explicit ``strptime`` format to try first.
        show_warnings: Whether to print extraction-failure warnings.

    Returns:
        Copy of the input DataFrame with an added ``date`` column (``NaT`` where
        extraction failed).

    Raises:
        ValueError: If ``filename_column`` is missing, or if month-only filenames
            were found and no ``default_year`` was supplied.
    """
    if filename_column not in df.columns:
        raise ValueError(
            f"Column '{filename_column}' not found in DataFrame. Available columns: "
            f"{list(df.columns)}"
        )

    df_copy = df.copy()

    month_only_patterns = [
        r"^(January|February|March|April|May|June|July|August|September|October|November"
        r"|December)(?:[^0-9]|$)",
        r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sept?|Oct|Nov|Dec)(?:[^0-9]|$)",
    ]

    unique_filenames = df_copy[filename_column].dropna().unique()
    month_only_found = False

    for filename in unique_filenames:
        filename_str = str(Path(filename).stem)
        for pattern in month_only_patterns:
            if re.search(pattern, filename_str, re.IGNORECASE) and not re.search(
                r"\d{4}", filename_str
            ):
                month_only_found = True
                break
        if month_only_found:
            break

    if month_only_found and default_year is None:
        raise ValueError(
            f"Found filenames with month-only patterns (no year) in column "
            f"'{filename_column}'.\nPlease provide a 'default_year' parameter "
            f"(e.g., default_year=2024) to handle these cases.\nExample filename: {filename}"
        )

    dates = []
    failed_extractions = []

    for _, row in df_copy.iterrows():
        filename = row[filename_column]

        if pd.isna(filename):
            dates.append(pd.NaT)
            continue

        extracted_date = extract_month_year_from_filename(filename, date_format=date_format)

        if extracted_date is None and default_year is not None:
            filename_str = str(Path(filename).stem)
            month_patterns = [
                (
                    r"(January|February|March|April|May|June|July|August|September|October"
                    r"|November|December)\.?",
                    "%B",
                ),
                (r"(Sept)\.?", "%b", "Sep"),
                (r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\.?", "%b"),
            ]

            for pattern_info in month_patterns:
                if len(pattern_info) == 3:
                    pattern, fmt, replacement = pattern_info
                else:
                    pattern, fmt = pattern_info
                    replacement = None

                match = re.search(pattern, filename_str, re.IGNORECASE)
                if match:
                    try:
                        month_str = match.group(1)
                        month_str = month_str.rstrip(".").capitalize()
                        if replacement:
                            month_str = replacement
                        date_str = f"{month_str} {default_year}"
                        date_obj = datetime.strptime(date_str, f"{fmt} %Y")
                        extracted_date = pd.Timestamp(date_obj)
                        break
                    except ValueError, AttributeError:
                        continue

        if extracted_date is None:
            failed_extractions.append(str(filename))
            dates.append(pd.NaT)
        else:
            dates.append(extracted_date)

    df_copy["date"] = dates

    if show_warnings and failed_extractions:
        print(f"\n⚠️  Warning: Could not extract dates from {len(failed_extractions)} filename(s):")
        for filename in failed_extractions[:5]:
            print(f"   - {filename}")
        if len(failed_extractions) > 5:
            print(f"   ... and {len(failed_extractions) - 5} more")
        print("\nThese rows will have NaT (Not a Time) in the 'date' column.")
        print("Consider providing a 'date_format' parameter if the format is non-standard.")

    if show_warnings:
        successful_extractions = df_copy["date"].notna().sum()
        total = len(df_copy)
        print(f"\n✓ Successfully extracted dates from {successful_extractions}/{total} rows")
        if successful_extractions > 0:
            unique_dates = df_copy["date"].dropna().unique()
            print(f"  Date range: {df_copy['date'].min()} to {df_copy['date'].max()}")
            print(f"  Unique dates found: {len(unique_dates)}")

    return df_copy


def experimental_excel_extraction_code(
    csv_filename: str | Path,
    OpenAI_client: Any,
    model: str = "gpt-4.1",
    desired_columns: list[str] | None = None,
    numeric_columns: list[str] | None = None,
    product_name_column: str = "Product name",
) -> pd.DataFrame:
    """Extract structured data from a CSV file using a large language model.

    Args:
        csv_filename: Path of the CSV file to process.
        OpenAI_client: Initialized OpenAI client used for the LLM call.
        model: Model identifier to use.
        desired_columns: Reserved for future use; currently ignored.
        numeric_columns: Reserved for future use; currently ignored.
        product_name_column: Reserved for future use; currently ignored.

    Returns:
        DataFrame parsed from the LLM-generated markdown table.
    """
    del desired_columns, numeric_columns, product_name_column

    with open(csv_filename, encoding="utf-8") as f:
        csv_data = f.read()

    markdown = experimental_csv_extraction_to_markdown(
        csv_data=csv_data,
        openai_client=OpenAI_client,
        model=model,
    )

    df = pd.read_csv(io.StringIO(markdown), sep="|")
    return df


def read_in_all_data_files(
    data_location: str | Path | None = None,
    file_type: Literal["csv", "excel"] | None = None,
    **read_kwargs: Any,
) -> tuple[list[pd.DataFrame], list[Path]]:
    """Read all CSV or Excel files from a directory into a list of DataFrames.

    Args:
        data_location: Directory to read files from.
        file_type: ``"csv"``, ``"excel"``, or ``None`` for auto-detection.
        **read_kwargs: Extra keyword arguments forwarded to the pandas reader.

    Returns:
        Tuple of (list of loaded DataFrames, list of source file paths).

    Raises:
        ValueError: If both file types are present without an explicit choice,
            no files are found, or ``file_type`` is invalid.
    """
    if data_location is None:
        raise ValueError("data_location is required and must point to a data directory")
    data_location = Path(data_location)

    if file_type is None:
        files = get_filtered_data_files(data_location, file_types=["csv", "excel"])
        csv_files = files["csv"]
        excel_files = files["excel"]

        if csv_files and not excel_files:
            file_type = "csv"
            print(f"Auto-detected file type: CSV ({len(csv_files)} files)")
        elif excel_files and not csv_files:
            file_type = "excel"
            print(f"Auto-detected file type: Excel ({len(excel_files)} files)")
        elif csv_files and excel_files:
            raise ValueError(
                f"Directory contains both CSV ({len(csv_files)}) and Excel "
                f"({len(excel_files)}) files. Please specify file_type='csv' or "
                "file_type='excel' explicitly."
            )
        else:
            raise ValueError(f"No CSV or Excel files found in {data_location}")

    if file_type == "csv":
        reader = pd.read_csv
    elif file_type == "excel":
        reader = pd.read_excel
    else:
        raise ValueError("file_type must be 'csv', 'excel', or None for auto-detection")

    files = get_filtered_data_files(data_location, file_types=[file_type])
    filepaths = files[file_type]
    dfs = [reader(fp, **read_kwargs) for fp in filepaths]
    return dfs, filepaths
