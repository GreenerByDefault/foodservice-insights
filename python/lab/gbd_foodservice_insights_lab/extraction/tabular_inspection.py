from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path
from typing import Any, Literal, TypedDict, overload

import openpyxl
import pandas as pd
from IPython.display import display

from gbd_foodservice_insights_lab.extraction.tabular_io import (
    _format_preview_cell,
    detect_excel_header_row,
    extract_excel_sheet_info,
    extract_month_year_from_filename,
    get_excel_headers,
    get_filtered_data_files,
    read_csv_with_auto_detection,
)

__all__ = [
    "check_same_columns",
    "compare_baseline_pilot_raw_data",
    "compare_component_datasets",
    "detect_component_duplicate_risk",
    "identify_low_cardinality_columns",
    "raw_data_report",
    "suggest_column_roles",
    "suggest_date_source_and_coverage",
    "suggest_tabular_import_settings",
    "validate_matching_columns",
]


class RawDataFileReport(TypedDict):
    """Files discovered by ``raw_data_report`` when details are requested."""

    pdf: list[Path]
    excel: list[Path]
    csv: list[Path]
    dirs: list[Path]
    single_excel_path: Path | None


class LowCardinalityDetails(TypedDict):
    """Evidence recorded for one low-cardinality column."""

    unique_count: int
    unique_values: Any
    null_count: int
    null_percentage: float
    dtype: str
    recommendation: str


class ColumnRoleCandidate(TypedDict):
    """One scored candidate for a suggested column role."""

    column: str
    score: int
    reason: str


class ColumnRoleDetails(TypedDict):
    """Best match and supporting candidates for one column role."""

    best_column: str | None
    confidence: str
    candidates: list[ColumnRoleCandidate]


class PairwiseOverlap(TypedDict):
    """Duplicate-row overlap between two component datasets."""

    left_component: str
    right_component: str
    intersection_count: int
    share_of_left: float
    share_of_right: float


def check_same_columns(df1: pd.DataFrame, df2: pd.DataFrame) -> None:
    """Compare column names of two DataFrames and print differences to the console.

    Args:
        df1: First DataFrame to compare.
        df2: Second DataFrame to compare.
    """
    cols_df1 = set(df1.columns)
    cols_df2 = set(df2.columns)

    if cols_df1 == cols_df2:
        print("The columns are the same.")
    else:
        print("The columns are not the same:")
        unique_to_df1 = cols_df1 - cols_df2
        if unique_to_df1:
            print(f"Columns unique to df1: {sorted(list(unique_to_df1))}")
        unique_to_df2 = cols_df2 - cols_df1
        if unique_to_df2:
            print(f"Columns unique to df2: {sorted(list(unique_to_df2))}")


def identify_low_cardinality_columns(
    df: pd.DataFrame,
    threshold: int = 10,
    show_details: bool = True,
    exclude_categorical: bool = False,
) -> list[Any]:
    """Identify columns with low cardinality and print analysis details.

    Args:
        df: DataFrame to inspect.
        threshold: Maximum unique-value count for a column to qualify.
        show_details: Whether to print per-column diagnostics.
        exclude_categorical: If ``True``, skip object-dtype columns.

    Returns:
        List of column names recommended for dropping.
    """
    low_cardinality_cols: dict[Any, LowCardinalityDetails] = {}

    for col in df.columns:
        if col in ["sheet_filename", "date"]:
            continue

        if exclude_categorical and df[col].dtype == "object":
            continue

        non_null_series = df[col].dropna()
        unique_vals = non_null_series.unique()
        unique_count = len(unique_vals)
        null_count = df[col].isnull().sum()

        if unique_count <= threshold:
            null_percentage = null_count / len(df) if len(df) else 0

            if unique_count == 1:
                recommendation = "DROP - Single value column"
            elif unique_count == 0:
                recommendation = "DROP - Empty column"
            elif null_percentage > 0.99:
                recommendation = "DROP - 99% null"
            elif null_percentage > 0.95:
                recommendation = "REVIEW - Mostly null values"
            elif df[col].dtype in ["object", "category"]:
                recommendation = "KEEP - Likely categorical variable"
            else:
                recommendation = "REVIEW - Low cardinality numeric"

            low_cardinality_cols[col] = {
                "unique_count": unique_count,
                "unique_values": unique_vals.tolist()
                if unique_count <= 20
                else f"First 20: {unique_vals[:20].tolist()}...",
                "null_count": null_count,
                "null_percentage": round(null_percentage * 100, 2),
                "dtype": str(df[col].dtype),
                "recommendation": recommendation,
            }

            if show_details:
                print(f"\nColumn: '{col}'")
                print(f"  Type: {df[col].dtype}")
                print(f"  Unique values: {unique_count}")
                print(f"  Null values: {null_count} ({null_count / len(df) * 100:.1f}%)")
                print(f"  Values: {low_cardinality_cols[col]['unique_values']}")
                print(f"  Recommendation: {recommendation}")

    if show_details and low_cardinality_cols:
        print(
            f"\nSummary: Found {len(low_cardinality_cols)} columns with <= {threshold} "
            "unique values"
        )

    print("returning columns that you should definitely drop. Consider dropping others")
    return [
        col for col, details in low_cardinality_cols.items() if "DROP" in details["recommendation"]
    ]


def validate_matching_columns(dfs: list[pd.DataFrame]) -> None:
    """Validate that all DataFrames in a list have exactly matching columns.

    Args:
        dfs: DataFrames to compare against the first one.

    Raises:
        ValueError: If any DataFrame has columns that do not match the first.
    """
    if len(dfs) < 2:
        return

    reference_cols = set(dfs[0].columns)

    for i, df in enumerate(dfs[1:], start=1):
        current_cols = set(df.columns)
        if current_cols != reference_cols:
            missing_in_current = reference_cols - current_cols
            extra_in_current = current_cols - reference_cols
            error_msg = f"Column mismatch in dataset {i}:\n"
            if missing_in_current:
                error_msg += f"  Missing columns: {sorted(missing_in_current)}\n"
            if extra_in_current:
                error_msg += f"  Extra columns: {sorted(extra_in_current)}\n"
            raise ValueError(error_msg)


def compare_baseline_pilot_raw_data(
    baseline_raw_path: str | Path,
    pilot_raw_path: str | Path,
) -> dict[str, Any]:
    """Compare raw baseline and pilot data to identify structural differences.

    Args:
        baseline_raw_path: Directory containing baseline raw data.
        pilot_raw_path: Directory containing pilot raw data.

    Returns:
        Dictionary describing file counts, column comparisons, row-count
        warnings, and a human-readable summary.
    """
    baseline_path = Path(baseline_raw_path)
    pilot_path = Path(pilot_raw_path)

    results = {
        "file_count_match": True,
        "baseline_counts": {},
        "pilot_counts": {},
        "baseline_columns": {},
        "pilot_columns": {},
        "column_mismatches": [],
        "row_count_warning": False,
        "summary": [],
    }

    baseline_files = get_filtered_data_files(baseline_path, file_types=["csv", "excel"])
    pilot_files = get_filtered_data_files(pilot_path, file_types=["csv", "excel"])

    baseline_csv_files = baseline_files["csv"]
    baseline_excel_files = baseline_files["excel"]
    pilot_csv_files = pilot_files["csv"]
    pilot_excel_files = pilot_files["excel"]

    results["baseline_counts"] = {
        "csv": len(baseline_csv_files),
        "excel": len(baseline_excel_files),
        "total": len(baseline_csv_files) + len(baseline_excel_files),
        "total_rows": 0,
    }
    results["pilot_counts"] = {
        "csv": len(pilot_csv_files),
        "excel": len(pilot_excel_files),
        "total": len(pilot_csv_files) + len(pilot_excel_files),
        "total_rows": 0,
    }

    if results["baseline_counts"]["csv"] != results["pilot_counts"]["csv"]:
        results["file_count_match"] = False
    if results["baseline_counts"]["excel"] != results["pilot_counts"]["excel"]:
        results["file_count_match"] = False

    def get_csv_info(filepath: Path) -> tuple[list, int]:
        try:
            df = pd.read_csv(filepath)
            return list(df.columns), len(df)
        except Exception as e:
            print(f"Error reading {filepath.name}: {e}")
            return [], 0

    results["baseline_columns"]["csv_files"] = {}
    for csv_file in baseline_csv_files:
        cols, row_count = get_csv_info(csv_file)
        results["baseline_columns"]["csv_files"][csv_file.name] = cols
        results["baseline_counts"]["total_rows"] += row_count

    results["pilot_columns"]["csv_files"] = {}
    for csv_file in pilot_csv_files:
        cols, row_count = get_csv_info(csv_file)
        results["pilot_columns"]["csv_files"][csv_file.name] = cols
        results["pilot_counts"]["total_rows"] += row_count

    results["baseline_columns"]["excel_files"] = {}
    baseline_excel_sheet_counts = []
    for excel_file in baseline_excel_files:
        sheet_cols, row_count = extract_excel_sheet_info(excel_file)
        results["baseline_columns"]["excel_files"][excel_file.name] = sheet_cols
        results["baseline_counts"]["total_rows"] += row_count
        baseline_excel_sheet_counts.append(len(sheet_cols))

    results["pilot_columns"]["excel_files"] = {}
    pilot_excel_sheet_counts = []
    for excel_file in pilot_excel_files:
        sheet_cols, row_count = extract_excel_sheet_info(excel_file)
        results["pilot_columns"]["excel_files"][excel_file.name] = sheet_cols
        results["pilot_counts"]["total_rows"] += row_count
        pilot_excel_sheet_counts.append(len(sheet_cols))

    baseline_csv_cols_list = list(results["baseline_columns"]["csv_files"].values())
    if baseline_csv_cols_list and not all(
        cols == baseline_csv_cols_list[0] for cols in baseline_csv_cols_list
    ):
        results["column_mismatches"].append(
            {
                "type": "baseline_csv_inconsistent",
                "message": "Not all baseline CSV files have the same columns",
            }
        )

    pilot_csv_cols_list = list(results["pilot_columns"]["csv_files"].values())
    if pilot_csv_cols_list and not all(
        cols == pilot_csv_cols_list[0] for cols in pilot_csv_cols_list
    ):
        results["column_mismatches"].append(
            {
                "type": "pilot_csv_inconsistent",
                "message": "Not all pilot CSV files have the same columns",
            }
        )

    if baseline_csv_cols_list and pilot_csv_cols_list:
        baseline_csv_cols = set(baseline_csv_cols_list[0])
        pilot_csv_cols = set(pilot_csv_cols_list[0])

        if baseline_csv_cols != pilot_csv_cols:
            cols_added = sorted(list(pilot_csv_cols - baseline_csv_cols))
            cols_removed = sorted(list(baseline_csv_cols - pilot_csv_cols))
            results["column_mismatches"].append(
                {
                    "type": "csv_columns_differ",
                    "message": "CSV files have different columns between baseline and pilot",
                    "columns_added": cols_added,
                    "columns_removed": cols_removed,
                    "order_changed": baseline_csv_cols_list[0] != pilot_csv_cols_list[0],
                }
            )

    if baseline_excel_sheet_counts and pilot_excel_sheet_counts:
        if not all(
            count == baseline_excel_sheet_counts[0] for count in baseline_excel_sheet_counts
        ):
            results["column_mismatches"].append(
                {
                    "type": "baseline_excel_sheet_count_inconsistent",
                    "message": (
                        "Baseline Excel files have varying sheet counts: "
                        f"{baseline_excel_sheet_counts}"
                    ),
                }
            )

        if not all(count == pilot_excel_sheet_counts[0] for count in pilot_excel_sheet_counts):
            results["column_mismatches"].append(
                {
                    "type": "pilot_excel_sheet_count_inconsistent",
                    "message": (
                        f"Pilot Excel files have varying sheet counts: {pilot_excel_sheet_counts}"
                    ),
                }
            )

        if baseline_excel_sheet_counts[0] != pilot_excel_sheet_counts[0]:
            results["column_mismatches"].append(
                {
                    "type": "excel_sheet_count_mismatch",
                    "message": (
                        f"Excel sheet count mismatch: baseline has "
                        f"{baseline_excel_sheet_counts[0]} sheets, pilot has "
                        f"{pilot_excel_sheet_counts[0]} sheets"
                    ),
                }
            )

    if baseline_excel_files and pilot_excel_files:
        baseline_sheet_columns_by_position = []
        for file_sheets in results["baseline_columns"]["excel_files"].values():
            sorted_sheets = sorted(file_sheets.items())
            for i, (_, cols) in enumerate(sorted_sheets):
                if i >= len(baseline_sheet_columns_by_position):
                    baseline_sheet_columns_by_position.append([])
                baseline_sheet_columns_by_position[i].append(tuple(cols))

        pilot_sheet_columns_by_position = []
        for file_sheets in results["pilot_columns"]["excel_files"].values():
            sorted_sheets = sorted(file_sheets.items())
            for i, (_, cols) in enumerate(sorted_sheets):
                if i >= len(pilot_sheet_columns_by_position):
                    pilot_sheet_columns_by_position.append([])
                pilot_sheet_columns_by_position[i].append(tuple(cols))

        for i in range(
            max(len(baseline_sheet_columns_by_position), len(pilot_sheet_columns_by_position))
        ):
            if i >= len(baseline_sheet_columns_by_position):
                results["column_mismatches"].append(
                    {
                        "type": "excel_extra_pilot_sheet",
                        "message": f"Pilot has extra sheet at position {i + 1}",
                    }
                )
                continue
            if i >= len(pilot_sheet_columns_by_position):
                results["column_mismatches"].append(
                    {
                        "type": "excel_missing_pilot_sheet",
                        "message": f"Pilot missing sheet at position {i + 1} (present in baseline)",
                    }
                )
                continue

            baseline_sheets_at_pos = baseline_sheet_columns_by_position[i]
            pilot_sheets_at_pos = pilot_sheet_columns_by_position[i]

            if baseline_sheets_at_pos and not all(
                cols == baseline_sheets_at_pos[0] for cols in baseline_sheets_at_pos
            ):
                results["column_mismatches"].append(
                    {
                        "type": "baseline_excel_inconsistent",
                        "message": (
                            f"Baseline Excel sheet at position {i + 1} has inconsistent "
                            "columns across files"
                        ),
                    }
                )

            if pilot_sheets_at_pos and not all(
                cols == pilot_sheets_at_pos[0] for cols in pilot_sheets_at_pos
            ):
                results["column_mismatches"].append(
                    {
                        "type": "pilot_excel_inconsistent",
                        "message": (
                            f"Pilot Excel sheet at position {i + 1} has inconsistent "
                            "columns across files"
                        ),
                    }
                )

            if baseline_sheets_at_pos and pilot_sheets_at_pos:
                baseline_cols = set(baseline_sheets_at_pos[0])
                pilot_cols = set(pilot_sheets_at_pos[0])

                if baseline_cols != pilot_cols:
                    cols_added = sorted(list(pilot_cols - baseline_cols))
                    cols_removed = sorted(list(baseline_cols - pilot_cols))
                    results["column_mismatches"].append(
                        {
                            "type": "excel_sheet_columns_differ",
                            "message": (
                                f"Excel sheet at position {i + 1} has different columns "
                                "between baseline and pilot"
                            ),
                            "sheet_position": i + 1,
                            "columns_added": cols_added,
                            "columns_removed": cols_removed,
                            "order_changed": list(baseline_sheets_at_pos[0])
                            != list(pilot_sheets_at_pos[0]),
                        }
                    )

    baseline_total_rows = results["baseline_counts"]["total_rows"]
    pilot_total_rows = results["pilot_counts"]["total_rows"]

    if baseline_total_rows > 0 and pilot_total_rows > 0:
        row_diff_percent = abs(baseline_total_rows - pilot_total_rows) / baseline_total_rows * 100
        if row_diff_percent > 20:
            results["row_count_warning"] = True
            results["column_mismatches"].append(
                {
                    "type": "row_count_warning",
                    "message": (
                        f"Total row count differs by {row_diff_percent:.1f}% "
                        f"(baseline: {baseline_total_rows:,}, pilot: {pilot_total_rows:,})"
                    ),
                    "severity": "warning",
                }
            )

    summary = []
    if not results["file_count_match"]:
        summary.append("⚠️  File count mismatch between baseline and pilot")
    if results["row_count_warning"]:
        summary.append("⚠️  Total row counts differ by >20%")
    if results["column_mismatches"]:
        non_warning_count = sum(
            1 for m in results["column_mismatches"] if m.get("severity") != "warning"
        )
        if non_warning_count > 0:
            summary.append(f"⚠️  {non_warning_count} column structure difference(s) found")

    if not summary:
        summary.append("✅ No structural differences found - pilot data matches baseline format")

    results["summary"] = summary

    print("=" * 80)
    print("BASELINE vs PILOT RAW DATA COMPARISON")
    print("=" * 80)
    print()

    print("📁 FILE COUNTS:")
    print(
        f"  Baseline: {results['baseline_counts']['csv']} CSV, "
        f"{results['baseline_counts']['excel']} Excel "
        f"(total: {results['baseline_counts']['total']})"
    )
    print(
        f"  Pilot:    {results['pilot_counts']['csv']} CSV, "
        f"{results['pilot_counts']['excel']} Excel "
        f"(total: {results['pilot_counts']['total']})"
    )
    if results["file_count_match"]:
        print("  ✅ File counts match")
    else:
        print("  ⚠️  File counts DO NOT match")
    print()

    print("📊 ROW COUNTS:")
    print(f"  Baseline: {results['baseline_counts']['total_rows']:,} total rows")
    print(f"  Pilot:    {results['pilot_counts']['total_rows']:,} total rows")
    if baseline_total_rows > 0 and pilot_total_rows > 0:
        row_diff_percent = abs(baseline_total_rows - pilot_total_rows) / baseline_total_rows * 100
        if row_diff_percent <= 20:
            print(f"  ✅ Row counts within 20% ({row_diff_percent:.1f}% difference)")
        else:
            print(f"  ⚠️  Row counts differ by {row_diff_percent:.1f}% (>20% threshold)")
    print()

    for msg in summary:
        print(msg)
    print()

    if results["column_mismatches"]:
        print("\n📋 DETAILED DIFFERENCES:")
        for i, mismatch in enumerate(results["column_mismatches"], 1):
            severity_marker = "⚠️ " if mismatch.get("severity") != "warning" else "ℹ️  "
            print(f"\n  {i}. {severity_marker}{mismatch['message']}")
            if mismatch.get("columns_added"):
                print(f"     + Columns added in pilot: {mismatch['columns_added']}")
            if mismatch.get("columns_removed"):
                print(f"     - Columns removed in pilot: {mismatch['columns_removed']}")
            if mismatch.get("order_changed"):
                print("     ⚠️  Column order changed")

    print("\n" + "=" * 80)
    return results


@overload
def raw_data_report(
    path: str | Path,
    recursive: bool = False,
    return_dict: Literal[False] = False,
    max_preview_rows: int = 10,
    baseline_pilot: str | None = None,
    base_filepath: str | None = None,
) -> str: ...


@overload
def raw_data_report(
    path: str | Path,
    recursive: bool = False,
    return_dict: Literal[True] = True,
    max_preview_rows: int = 10,
    baseline_pilot: str | None = None,
    base_filepath: str | None = None,
) -> tuple[str, RawDataFileReport]: ...


@overload
def raw_data_report(
    path: str | Path,
    recursive: bool = False,
    return_dict: bool = False,
    max_preview_rows: int = 10,
    baseline_pilot: str | None = None,
    base_filepath: str | None = None,
) -> str | tuple[str, RawDataFileReport]: ...


def raw_data_report(
    path: str | Path,
    recursive: bool = False,
    return_dict: bool = False,
    max_preview_rows: int = 10,
    baseline_pilot: str | None = None,
    base_filepath: str | None = None,
) -> str | tuple[str, RawDataFileReport]:
    """Print a raw-data report and return a data type classification.

    Args:
        path: Directory of raw data files to inspect.
        recursive: Whether to recurse into subdirectories.
        return_dict: If ``True``, also return the discovered file paths.
        max_preview_rows: Number of preview rows to show per file.
        baseline_pilot: Optional ``"baseline"`` / ``"pilot"`` marker for
            triggering a baseline-vs-pilot comparison.
        base_filepath: Filepath used to locate the matching dataset for the
            baseline-vs-pilot comparison.

    Returns:
        Either the data type string, or a tuple of (data type, dictionary of
        discovered file paths) when ``return_dict`` is ``True``.
    """
    base = Path(path)

    files = get_filtered_data_files(base, recursive=recursive)
    pdf_files = files["pdf"]
    excel_files = files["excel"]
    csv_files = files["csv"]

    subdirs = sorted({p.parent for p in pdf_files + excel_files + csv_files if p.parent != base})

    data_type = "other"
    single_excel_path: Path | None = None

    print("=" * 80)
    print("RAW DATA REPORT".center(80))
    print("=" * 80)
    print(f"\n📁 Directory: {base}")
    print(f"🔍 Scan mode: {'Recursive' if recursive else 'Non-recursive'}\n")

    summary_data = {
        "File Type": ["PDF Files", "Excel Files", "CSV Files", "Total Spreadsheets"],
        "Count": [
            len(pdf_files),
            len(excel_files),
            len(csv_files),
            len(excel_files) + len(csv_files),
        ],
    }
    summary_df = pd.DataFrame(summary_data)
    print("📊 SUMMARY:")
    display(summary_df)

    if subdirs:
        print(f"\n📂 Subdirectories found: {len(subdirs)}")
        for d in subdirs:
            print(f"   └─ {d.relative_to(base)}")

    print("\n" + "-" * 80 + "\n")

    if pdf_files:
        print("📄 PDF FILES:")
        for i, pf in enumerate(pdf_files, 1):
            print(f"   {i}. {pf.relative_to(base)}")
        print()

    excel_headers = []
    if excel_files:
        print("=" * 80)
        print("📗 EXCEL FILE DETAILS")
        print("=" * 80)
        print("\n⚠️  NOTE: Row counts shown below use Excel's max_row property, which may")
        print("   include empty rows with formatting. Pandas may read fewer actual data rows.")
        print("   Use compare_component_datasets() to see true loaded row counts.\n")

        for file_idx, xf in enumerate(excel_files, 1):
            try:
                wb = openpyxl.load_workbook(xf, read_only=False, data_only=False)

                print(f"[{file_idx}/{len(excel_files)}] {xf.relative_to(base)}")
                print(f"{'─' * 80}")

                for sheet_idx, ws in enumerate(wb.worksheets, 1):
                    header_row = detect_excel_header_row(ws)
                    hdr = get_excel_headers(ws, header_row=header_row)
                    excel_headers.append(tuple(hdr))

                    data_rows = max(ws.max_row - header_row, 0)

                    if len(wb.worksheets) > 1:
                        print(f"\n  📋 Sheet {sheet_idx}: '{ws.title}'")
                    print(f"  📊 Dimensions: {data_rows:,} rows × {ws.max_column} columns")
                    print(f"  🧠 Likely header row: {header_row}")

                    print(f"\n  📝 Columns ({len(hdr)}):")
                    for i, col in enumerate(hdr, 1):
                        if i % 3 == 1:
                            print("     ", end="")
                        print(f"{_format_preview_cell(col):<30}", end="")
                        if i % 3 == 0 or i == len(hdr):
                            print()
                    if len(hdr) % 3 != 0:
                        print()

                    if max_preview_rows > 0 and data_rows > 0:
                        print(
                            f"\n  🔍 Data Preview (first {min(max_preview_rows, data_rows)} rows):"
                        )
                        rows_data = []
                        preview_end_row = min(header_row + max_preview_rows, ws.max_row)
                        for row in ws.iter_rows(
                            min_row=header_row + 1,
                            max_row=preview_end_row,
                            values_only=True,
                        ):
                            rows_data.append(list(row))

                        if rows_data:
                            preview_df = pd.DataFrame(rows_data, columns=hdr)
                            display(preview_df)

                    if sheet_idx < len(wb.worksheets):
                        print(f"\n  {'-' * 76}")

                print("\n")

            except Exception as e:
                print(f"❌ ERROR reading {xf.relative_to(base)}: {e}\n")

    csv_headers = []
    if csv_files:
        print("=" * 80)
        print("📘 CSV FILE DETAILS")
        print("=" * 80 + "\n")

        for file_idx, cf in enumerate(csv_files, 1):
            try:
                df_preview, encoding, delimiter = read_csv_with_auto_detection(
                    cf,
                    nrows=max_preview_rows,
                )
                hdr = df_preview.columns.tolist()
                csv_headers.append(tuple(hdr))

                with open(cf, newline="", encoding=encoding) as f:
                    nrows = sum(1 for _ in csv.reader(f, delimiter=delimiter)) - 1

                print(f"[{file_idx}/{len(csv_files)}] {cf.relative_to(base)}")
                print(f"{'─' * 80}")
                print(f"  📊 Dimensions: {nrows:,} rows × {len(hdr)} columns")
                print(f"  🔤 Encoding: {encoding}")
                print(f"  🔸 Delimiter: {delimiter!r}")

                print(f"\n  📝 Columns ({len(hdr)}):")
                for i, col in enumerate(hdr, 1):
                    if i % 3 == 1:
                        print("     ", end="")
                    print(f"{col:<30}", end="")
                    if i % 3 == 0 or i == len(hdr):
                        print()
                if len(hdr) % 3 != 0:
                    print()

                if max_preview_rows > 0 and len(df_preview) > 0:
                    print(f"\n  🔍 Data Preview (first {len(df_preview)} rows):")
                    display(df_preview)

                print("\n")

            except Exception as e:
                print(f"❌ ERROR reading {cf.relative_to(base)}: {e}\n")

    all_headers = excel_headers + csv_headers
    if all_headers:
        all_unique_cols = set()
        for header_tuple in all_headers:
            all_unique_cols.update(header_tuple)

        date_keywords = ["date", "month"]
        possible_date_cols = [
            col
            for col in all_unique_cols
            if any(keyword in str(col).lower() for keyword in date_keywords)
        ]

        if possible_date_cols:
            print("=" * 80)
            print("📅 DATE COLUMN DETECTION")
            print("=" * 80 + "\n")
            print(f"✅ Possible date column(s) detected: {sorted(possible_date_cols)}")
            print("=" * 80 + "\n")

    if len(all_headers) > 1:
        print("=" * 80)
        print("🔍 COLUMN CONSISTENCY CHECK")
        print("=" * 80 + "\n")

        unique_headers = set(all_headers)
        if len(unique_headers) == 1:
            print("✅ All Excel and CSV files have IDENTICAL columns!")
            print(f"   → {len(all_headers)} file(s) with {len(all_headers[0])} columns each")
        else:
            print("⚠️  WARNING: Files have DIFFERENT column structures!")
            print(
                f"   → Found {len(unique_headers)} different column configurations "
                f"across {len(all_headers)} files\n"
            )

            all_cols = [col for h in all_headers for col in h]
            col_counts = Counter(all_cols)
            total_files = len(all_headers)

            inconsistent = [col for col, count in col_counts.items() if count < total_files]
            if inconsistent:
                print(f"❌ Columns NOT present in all files ({len(inconsistent)} columns):")
                inconsistent_data = []
                for col in sorted(inconsistent):
                    count = col_counts[col]
                    inconsistent_data.append(
                        {
                            "Column Name": col,
                            "Present In": f"{count}/{total_files} files",
                            "% Coverage": f"{(count / total_files) * 100:.0f}%",
                        }
                    )
                inconsistent_df = pd.DataFrame(inconsistent_data)
                display(inconsistent_df)
                print()

            print("📋 Column Configuration Summary:\n")
            for i, header_set in enumerate(unique_headers, 1):
                matching_files = sum(1 for h in all_headers if h == header_set)
                print(f"   Configuration {i}:")
                print(f"   └─ {matching_files} file(s) with {len(header_set)} columns")
                if len(header_set) <= 10:
                    print(f"   └─ Columns: {list(header_set)}")
                else:
                    print(f"   └─ First 10 columns: {list(header_set)[:10]}...")
                print()

        print("=" * 80)

    if len(excel_files) == 1 and len(csv_files) == 0:
        try:
            wb = openpyxl.load_workbook(excel_files[0], read_only=True)
            sheet_count = len(wb.sheetnames)
            wb.close()
            if sheet_count >= 2:
                data_type = "single_multitab_excel"
                single_excel_path = excel_files[0]
            else:
                data_type = "single_tab_files"
        except Exception:
            data_type = "other"
    elif len(csv_files) > 0 and len(excel_files) == 0:
        data_type = "single_tab_files"
    elif len(excel_files) > 0 and len(csv_files) == 0:
        all_single_sheet = True
        for xf in excel_files:
            try:
                wb = openpyxl.load_workbook(xf, read_only=True)
                if len(wb.sheetnames) > 1:
                    all_single_sheet = False
                wb.close()
            except Exception:
                all_single_sheet = False
                break
        data_type = "single_tab_files" if all_single_sheet else "other"
    else:
        data_type = "other"

    print("\n" + "=" * 80)
    print("📋 DATA TYPE CLASSIFICATION")
    print("=" * 80)
    print(f"\n🏷️  Detected type: {data_type}")
    if data_type == "single_multitab_excel":
        print(
            "   → Single Excel file with multiple sheets: "
            f"{single_excel_path.name if single_excel_path else 'N/A'}"
        )
    elif data_type == "single_tab_files":
        file_type = "CSV" if csv_files else "Excel"
        file_count = len(csv_files) if csv_files else len(excel_files)
        if file_count == 1:
            print(f"   → Single {file_type} file")
        else:
            print(f"   → Multiple {file_type} files to concatenate")
    else:
        print("   → Manual inspection required - mixed file types or complex structure")
    print("=" * 80 + "\n")

    if baseline_pilot == "pilot" and base_filepath is not None:
        baseline_raw_path = base_filepath.replace("pilot", "baseline") + "/raw_data"
        compare_baseline_pilot_raw_data(baseline_raw_path, path)

    if return_dict:
        return data_type, {
            "pdf": pdf_files,
            "excel": excel_files,
            "csv": csv_files,
            "dirs": subdirs,
            "single_excel_path": single_excel_path,
        }

    return data_type


def compare_component_datasets(dfs: list[pd.DataFrame]) -> None:
    """Compare a list of DataFrames for consistency in row counts and column names.

    Args:
        dfs: Component DataFrames to compare.
    """
    row_nums: list[int] = [df.shape[0] for df in dfs]
    print(f"comparing number of rows:  {row_nums}")

    if len(row_nums) > 1:
        max_rows = max(row_nums)
        min_rows = min(row_nums)
        if min_rows > 0:
            pct_difference = ((max_rows - min_rows) / min_rows) * 100
            if pct_difference > 30:
                print(
                    "⚠️  WARNING: Row counts vary significantly! Largest dataset has "
                    f"{pct_difference:.1f}% more rows than smallest."
                )
    print("-" * 30, "\n")

    if not dfs:
        print("Warning: No DataFrames provided to compare.")
        return

    if len(dfs) < 2:
        print("Comparison requires at least two DataFrames.\n")
        return

    col_sets = [set(df.columns) for df in dfs]
    num_dfs = len(dfs)

    common_cols = set.intersection(*col_sets)
    if common_cols:
        print(f"Columns found in all {num_dfs} DataFrames:   {sorted(list(common_cols))}")
    else:
        print("No columns common to all DataFrames.")
    print("-" * 30, "\n")

    print("Columns unique to a single DataFrame:")
    all_unique_cols = set()
    found_any_unique = False
    for i, current_set in enumerate(col_sets):
        other_sets = col_sets[:i] + col_sets[i + 1 :]
        union_of_others = set.union(*other_sets) if other_sets else set()

        unique_to_current = current_set - union_of_others
        if unique_to_current:
            found_any_unique = True
            print(f"  - In DataFrame {i}: {sorted(list(unique_to_current))}")
            all_unique_cols.update(unique_to_current)

    if not found_any_unique:
        print("  None")
    print("-" * 30, "\n")

    all_cols_union = set.union(*col_sets)
    partially_shared_cols = all_cols_union - common_cols - all_unique_cols

    print("Columns found in multiple (but not all) DataFrames:")
    if partially_shared_cols:
        print(f"  {sorted(list(partially_shared_cols))}")
    else:
        print("  None")


def _component_names_from_inputs(
    dfs: list[pd.DataFrame],
    component_names: list[str] | None = None,
) -> list[str]:
    if component_names is not None:
        if len(component_names) != len(dfs):
            raise ValueError(
                f"dfs and component_names must have the same length; got {len(dfs)} "
                f"DataFrames and {len(component_names)} names"
            )
        return [str(name) for name in component_names]
    return [f"component_{i + 1}" for i in range(len(dfs))]


def suggest_tabular_import_settings(
    data_location: str | Path,
    recursive: bool = False,
    scan_rows: int = 20,
) -> dict[str, Any]:
    """Suggest likely header rows and skiprows values for tabular imports.

    Args:
        data_location: Directory of CSV/Excel files to inspect.
        recursive: Whether to recurse into subdirectories.
        scan_rows: Number of leading rows to consider when scoring headers.

    Returns:
        Dictionary describing per-component suggestions and a recommended shared
        ``skiprows`` value when one exists.
    """
    base = Path(data_location)
    files = get_filtered_data_files(base, file_types=["csv", "excel"], recursive=recursive)
    components = []

    for csv_file in files["csv"]:
        rel_name = str(csv_file.relative_to(base))
        components.append(
            {
                "component_name": rel_name,
                "file_type": "csv",
                "header_row": 1,
                "suggested_skiprows": 0,
            }
        )

    for excel_file in files["excel"]:
        wb = openpyxl.load_workbook(excel_file, read_only=True)
        multi_sheet = len(wb.sheetnames) > 1
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_row = detect_excel_header_row(ws, scan_rows=scan_rows)
            component_name = str(excel_file.relative_to(base))
            if multi_sheet:
                component_name = f"{component_name}::{sheet_name}"
            components.append(
                {
                    "component_name": component_name,
                    "file_type": "excel",
                    "header_row": header_row,
                    "suggested_skiprows": max(header_row - 1, 0),
                }
            )
        wb.close()

    excel_skiprows = {
        component["suggested_skiprows"]
        for component in components
        if component["file_type"] == "excel"
    }

    if not components:
        recommended_skiprows = None
    elif excel_skiprows and len(excel_skiprows) == 1:
        recommended_skiprows = next(iter(excel_skiprows))
    elif all(component["file_type"] == "csv" for component in components):
        recommended_skiprows = 0
    else:
        recommended_skiprows = None

    result = {
        "components": components,
        "recommended_skiprows": recommended_skiprows,
        "has_consistent_excel_skiprows": len(excel_skiprows) <= 1,
    }

    print("=" * 80)
    print("🧭 IMPORT SETTINGS SUGGESTIONS")
    print("=" * 80)
    if not components:
        print("No CSV or Excel files found to inspect.")
    else:
        for component in components:
            print(
                f"{component['component_name']}: header row {component['header_row']}, "
                f"suggested skiprows={component['suggested_skiprows']}"
            )
        if recommended_skiprows is not None:
            print(f"\nRecommended shared skiprows value: {recommended_skiprows}")
        else:
            print("\nNo single shared skiprows value recommended across all components.")
    print("=" * 80 + "\n")
    return result


def _combined_sample_for_roles(dfs: list[pd.DataFrame]) -> pd.DataFrame:
    if not dfs:
        return pd.DataFrame()
    sample_frames = [df.head(100) for df in dfs]
    return pd.concat(sample_frames, ignore_index=True, sort=False)


def _keyword_hits(column_name: str, keywords: list[str]) -> list[str]:
    lower_name = column_name.lower()
    return [keyword for keyword in keywords if keyword in lower_name]


def _parse_success_rate(series: pd.Series) -> float:
    non_null = series.dropna().head(50)
    if non_null.empty:
        return 0.0
    try:
        parsed = pd.to_datetime(non_null, errors="coerce")
    except Exception:
        return 0.0
    return float(parsed.notna().mean())


def _numeric_success_rate(series: pd.Series) -> float:
    non_null = series.dropna().head(50)
    if non_null.empty:
        return 0.0
    converted = pd.to_numeric(non_null, errors="coerce")
    return float(converted.notna().mean())


def _score_role(series: pd.Series, column_name: str, role: str) -> tuple[int, list[str]]:
    lower_name = column_name.lower()
    score = 0
    reasons = []
    numeric_rate = _numeric_success_rate(series)
    parse_rate = _parse_success_rate(series)
    string_like = pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series)
    non_null = series.dropna()
    unique_ratio = float(non_null.nunique() / len(non_null)) if len(non_null) else 0.0
    avg_length = (
        float(non_null.astype(str).str.len().mean()) if len(non_null) and string_like else 0.0
    )

    role_keywords = {
        "product": ["product", "description", "item", "menu", "food", "name"],
        "weight": ["weight", "wght", "lbs", "lb", "kg", "kilo", "shipped", "net"],
        "date": ["date", "month", "period", "week"],
        "cases": ["case", "cases", "cs"],
        "pack_size": ["pack", "size", "uom", "unit", "units"],
        "spend": ["dollar", "dollars", "sales", "amount", "cost", "price", "spend", "avg$"],
    }

    hits = _keyword_hits(lower_name, role_keywords[role])
    if hits:
        score += min(len(hits) * 2, 4)
        reasons.append(f"name suggests {role}: {hits}")

    if role == "product":
        if string_like:
            score += 2
            reasons.append("mostly text-like values")
        if unique_ratio > 0.5:
            score += 1
            reasons.append("many unique values")
        if avg_length > 12:
            score += 1
            reasons.append("longer text values")
        if numeric_rate > 0.8:
            score -= 2
    elif role in {"weight", "cases", "spend"}:
        if numeric_rate > 0.8:
            score += 2
            reasons.append("mostly numeric values")
        if parse_rate > 0.8:
            score -= 2
        if role == "weight" and any(token in lower_name for token in ["qty", "quantity"]):
            score += 1
            reasons.append("name suggests quantity or shipped amount")
    elif role == "date":
        if parse_rate > 0.8:
            score += 3
            reasons.append("sample values parse as dates")
        elif parse_rate > 0.4:
            score += 1
            reasons.append("some values parse as dates")
        if numeric_rate > 0.8 and not hits:
            score -= 1
    elif role == "pack_size":
        if string_like:
            score += 1
            reasons.append("string-like unit text")
        if numeric_rate > 0.8:
            score -= 1

    return score, reasons


def _confidence_from_candidates(candidates: list[ColumnRoleCandidate]) -> str:
    if not candidates:
        return "low"
    best_score = candidates[0]["score"]
    second_score = candidates[1]["score"] if len(candidates) > 1 else -999
    if best_score >= 5 and best_score - second_score >= 2:
        return "high"
    if best_score >= 2:
        return "medium"
    return "low"


def suggest_column_roles(
    dfs: list[pd.DataFrame],
    component_names: list[str] | None = None,
) -> dict[str, Any]:
    """Suggest likely product, weight, date, and other key tabular columns.

    Args:
        dfs: Component DataFrames to score columns over.
        component_names: Optional friendly names for each component (currently
            unused but accepted for API consistency).

    Returns:
        Dictionary mapping each role to its best candidate column, confidence,
        and top scoring candidates.
    """
    del component_names

    sample_df = _combined_sample_for_roles(dfs)
    if sample_df.empty:
        result = {"roles": {}}
        print("No data available for column-role suggestions.")
        return result

    roles = ["product", "weight", "date", "cases", "pack_size", "spend"]
    role_results: dict[str, ColumnRoleDetails] = {}

    for role in roles:
        candidates: list[ColumnRoleCandidate] = []
        for column_label in sample_df.columns:
            column_name = str(column_label)
            score, reasons = _score_role(sample_df[column_label], column_name, role)
            candidates.append(
                {
                    "column": column_name,
                    "score": score,
                    "reason": "; ".join(reasons) if reasons else "weak signal",
                }
            )
        candidates.sort(key=lambda item: (-item["score"], item["column"]))
        role_results[role] = {
            "best_column": candidates[0]["column"] if candidates else None,
            "confidence": _confidence_from_candidates(candidates),
            "candidates": candidates[:3],
        }

    print("=" * 80)
    print("🧠 COLUMN ROLE SUGGESTIONS")
    print("=" * 80)
    for role, details in role_results.items():
        print(f"{role}: {details['best_column']} (confidence: {details['confidence']})")
        if details["candidates"]:
            print(f"  top reason: {details['candidates'][0]['reason']}")
    print("=" * 80 + "\n")

    return {"roles": role_results}


def _monthly_period_summary(period_strings: list[str]) -> dict[str, Any]:
    counts = Counter(period_strings)
    observed = sorted(counts)
    duplicate_periods = [period for period, count in counts.items() if count > 1]

    missing_periods: list[str] = []
    if observed:
        period_index = pd.PeriodIndex(observed, freq="M")
        full_range = pd.period_range(period_index.min(), period_index.max(), freq="M")
        missing_periods = [str(period) for period in full_range if str(period) not in counts]

    return {
        "observed_periods": observed,
        "duplicate_periods": sorted(duplicate_periods),
        "missing_periods": missing_periods,
    }


def suggest_date_source_and_coverage(
    dfs: list[pd.DataFrame],
    component_names: list[str] | None = None,
) -> dict[str, Any]:
    """Suggest the most likely date source and summarize observed period coverage.

    Args:
        dfs: Component DataFrames to inspect.
        component_names: Optional friendly names for each component, used when
            falling back to filename-based date inference.

    Returns:
        Dictionary describing the chosen date source, confidence, observed
        periods, missing periods, duplicates, and unparsed components.

    Raises:
        ValueError: If ``dfs`` and ``component_names`` have different lengths.
    """
    component_names = _component_names_from_inputs(dfs, component_names)
    sample_df = _combined_sample_for_roles(dfs)

    best_date_column = None
    best_date_score = -1.0
    date_column_candidates = []

    if not sample_df.empty:
        for column_name in sample_df.columns:
            score, reasons = _score_role(sample_df[column_name], column_name, "date")
            if score > best_date_score:
                best_date_score = score
                best_date_column = column_name
            date_column_candidates.append((column_name, score, reasons))

    if best_date_column is not None and best_date_score >= 3:
        component_period_rows = []
        period_strings: list[str] = []
        unparsed_components: list[str] = []

        for component_name, df in zip(component_names, dfs, strict=True):
            if best_date_column not in df.columns:
                unparsed_components.append(component_name)
                continue
            parsed = pd.to_datetime(df[best_date_column], errors="coerce")
            periods = sorted(parsed.dropna().dt.to_period("M").astype(str).unique().tolist())
            if periods:
                period_strings.extend(periods)
            else:
                unparsed_components.append(component_name)
            component_period_rows.append(
                {
                    "component_name": component_name,
                    "periods": periods,
                }
            )

        coverage = _monthly_period_summary(period_strings)
        result = {
            "source": "existing_column",
            "confidence": "high" if best_date_score >= 5 else "medium",
            "date_column": best_date_column,
            "component_periods": component_period_rows,
            "unparsed_components": unparsed_components,
            **coverage,
        }
    else:
        component_period_rows = []
        period_strings = []
        unparsed_components = []

        for component_name in component_names:
            parsed = extract_month_year_from_filename(component_name)
            if parsed is None:
                unparsed_components.append(component_name)
                component_period_rows.append({"component_name": component_name, "period": None})
                continue
            period_str = str(parsed.to_period("M"))
            period_strings.append(period_str)
            component_period_rows.append({"component_name": component_name, "period": period_str})

        coverage = _monthly_period_summary(period_strings)
        if period_strings:
            result = {
                "source": "component_name",
                "confidence": "medium",
                "date_column": None,
                "component_periods": component_period_rows,
                "unparsed_components": unparsed_components,
                **coverage,
            }
        else:
            result = {
                "source": "none",
                "confidence": "low",
                "date_column": None,
                "component_periods": component_period_rows,
                "unparsed_components": unparsed_components,
                **coverage,
            }

    print("=" * 80)
    print("📅 DATE SOURCE AND COVERAGE SUGGESTIONS")
    print("=" * 80)
    print(f"Likely date source: {result['source']} (confidence: {result['confidence']})")
    if result["date_column"]:
        print(f"Suggested date column: {result['date_column']}")
    if result["observed_periods"]:
        print(f"Observed periods: {result['observed_periods']}")
    if result["missing_periods"]:
        print(f"Missing periods inside observed range: {result['missing_periods']}")
    if result["duplicate_periods"]:
        print(f"Repeated periods: {result['duplicate_periods']}")
    if result["unparsed_components"]:
        print(f"Unparsed components: {result['unparsed_components']}")
    print("=" * 80 + "\n")

    return result


def _normalize_duplicate_value(value: Any) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{float(value):.12g}"

    text = str(value).strip()
    if not text:
        return ""

    try:
        numeric_value = float(text.replace(",", ""))
        return f"{numeric_value:.12g}"
    except ValueError:
        pass

    parsed = pd.to_datetime(pd.Series([text]), errors="coerce")
    if parsed.notna().iloc[0]:
        return parsed.iloc[0].isoformat()

    return " ".join(text.lower().split())


def detect_component_duplicate_risk(
    dfs: list[pd.DataFrame],
    component_names: list[str] | None = None,
    duplicate_share_threshold: float = 0.20,
) -> dict[str, Any]:
    """Flag exact duplicate risk across component datasets before Step 1.

    Args:
        dfs: Component DataFrames to compare.
        component_names: Optional friendly names for each component.
        duplicate_share_threshold: Share of duplicate rows that flips the status
            to ``"warning"``.

    Returns:
        Dictionary describing duplicate counts, per-source contributions, and
        pairwise overlap metrics.

    Raises:
        ValueError: If ``dfs`` and ``component_names`` have different lengths.
    """
    component_names = _component_names_from_inputs(dfs, component_names)

    if not dfs:
        return {
            "status": "not_checked",
            "threshold_crossed": False,
            "duplicate_row_count": 0,
            "duplicate_row_share": 0.0,
            "duplicate_pattern_count": 0,
            "compare_columns": [],
            "per_source_contribution": [],
            "pairwise_overlap": [],
            "message": "No component datasets supplied.",
        }

    common_cols = set.intersection(*(set(df.columns) for df in dfs))
    compare_cols = sorted(
        col
        for col in common_cols
        if col not in {"sheet_filename", "original_sheet_name", "source_component"}
    )

    if not compare_cols:
        result = {
            "status": "not_checked",
            "threshold_crossed": False,
            "duplicate_row_count": 0,
            "duplicate_row_share": 0.0,
            "duplicate_pattern_count": 0,
            "compare_columns": [],
            "per_source_contribution": [],
            "pairwise_overlap": [],
            "message": "No shared business columns were available for duplicate checking.",
        }
        print(result["message"])
        return result

    prepared_frames = []
    for component_name, df in zip(component_names, dfs, strict=True):
        subset = df.loc[:, compare_cols].copy()
        subset = subset.dropna(how="all")
        for column in compare_cols:
            subset[column] = subset[column].map(_normalize_duplicate_value)
        subset = subset.loc[~(subset.eq("").all(axis=1))].copy()
        subset["source_component"] = component_name
        prepared_frames.append(subset)

    combined = pd.concat(prepared_frames, ignore_index=True)
    duplicate_mask = combined.duplicated(subset=compare_cols, keep=False)
    duplicate_rows = combined.loc[duplicate_mask].copy()

    duplicate_row_count = int(duplicate_mask.sum())
    duplicate_row_share = float(duplicate_row_count / len(combined)) if len(combined) else 0.0
    duplicate_pattern_count = (
        int(duplicate_rows.groupby(compare_cols, dropna=False).ngroups)
        if not duplicate_rows.empty
        else 0
    )

    per_source_contribution = []
    for component_name, prepared in zip(component_names, prepared_frames, strict=True):
        source_rows = duplicate_rows.loc[duplicate_rows["source_component"] == component_name]
        per_source_contribution.append(
            {
                "component_name": component_name,
                "row_count": len(prepared),
                "duplicate_rows": len(source_rows),
                "duplicate_share_within_source": (
                    float(len(source_rows) / len(prepared)) if len(prepared) else 0.0
                ),
            }
        )

    pairwise_overlap: list[PairwiseOverlap] = []
    unique_rows_by_source = {}
    for component_name, prepared in zip(component_names, prepared_frames, strict=True):
        unique_rows = {
            tuple(row)
            for row in prepared.loc[:, compare_cols]
            .drop_duplicates()
            .itertuples(index=False, name=None)
        }
        unique_rows_by_source[component_name] = unique_rows

    for i, left_name in enumerate(component_names):
        for right_name in component_names[i + 1 :]:
            left_rows = unique_rows_by_source[left_name]
            right_rows = unique_rows_by_source[right_name]
            if not left_rows and not right_rows:
                continue
            intersection_count = len(left_rows & right_rows)
            pairwise_overlap.append(
                {
                    "left_component": left_name,
                    "right_component": right_name,
                    "intersection_count": intersection_count,
                    "share_of_left": float(intersection_count / len(left_rows))
                    if left_rows
                    else 0.0,
                    "share_of_right": float(intersection_count / len(right_rows))
                    if right_rows
                    else 0.0,
                }
            )

    pairwise_overlap.sort(
        key=lambda item: (
            -item["intersection_count"],
            -item["share_of_left"],
            -item["share_of_right"],
            item["left_component"],
            item["right_component"],
        )
    )

    threshold_crossed = duplicate_row_share >= duplicate_share_threshold
    status = "warning" if threshold_crossed else "ok"
    message = (
        (
            f"Duplicate risk is high: {duplicate_row_count} of {len(combined)} combined rows "
            f"({duplicate_row_share:.1%}) are exact duplicates across shared business columns."
        )
        if threshold_crossed
        else (
            f"Duplicate risk is below threshold: {duplicate_row_count} of {len(combined)} "
            f"combined rows ({duplicate_row_share:.1%}) are exact duplicates across shared "
            "business columns."
        )
    )

    result = {
        "status": status,
        "threshold_crossed": threshold_crossed,
        "duplicate_row_count": duplicate_row_count,
        "duplicate_row_share": duplicate_row_share,
        "duplicate_pattern_count": duplicate_pattern_count,
        "compare_columns": compare_cols,
        "per_source_contribution": per_source_contribution,
        "pairwise_overlap": pairwise_overlap,
        "message": message,
    }

    print("=" * 80)
    print("🔁 COMPONENT DUPLICATE-RISK CHECK")
    print("=" * 80)
    print(message)
    if pairwise_overlap:
        top_overlap = pairwise_overlap[0]
        print(
            f"Top overlap pair: {top_overlap['left_component']} vs "
            f"{top_overlap['right_component']} ({top_overlap['intersection_count']} shared "
            f"unique rows, {top_overlap['share_of_left']:.1%} of left, "
            f"{top_overlap['share_of_right']:.1%} of right)"
        )
    print("=" * 80 + "\n")

    return result
